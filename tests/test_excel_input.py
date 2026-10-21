from __future__ import annotations

from datetime import datetime

from openpyxl import Workbook, load_workbook

from rkn010_migration.excel_input import read_workbook


HEADERS = [
    "row_uuid",
    "Идентификатор записи",
    "Регистрационный номер",
    "ОГРН",
    "Полное наименование",
    "Сокращенное наименование",
    "Место нахождения",
    "N1",
    "Дата начала и окончания лицензии_1",
    "Статус_1",
    "N2",
    "Дата окончания лицензии_2",
    "Статус_2",
    "N3",
    "Дата начала и окончания лицензии_3",
    "Статус_3",
    "Зона нумерации",
    "Номер приказа о включении",
    "Дата приказа о включении",
    "Номер приказа об исключении",
    "Дата приказа об исключении",
    "Дата экспертного заключения",
]


def make_book(path, *, zone="843"):
    wb = Workbook()
    ws = wb.active
    ws.title = "РОСЗСП"
    ws.append(["Таблица миграции"])
    ws.append([])
    ws.append(HEADERS)
    ws.append(
        [
            "uuid",
            "legacy",
            "00000052",
            "1021600000001",
            "ПАО ТЕСТОВЫЙ ОПЕРАТОР",
            "ПАО ТЕСТ",
            "г. Казань",
            "3004",
            "01.01.2030",
            "действующая",
            "3004",
            "01.06.2029",
            "действующая",
            "9913",
            "действующая",
            "01.01.2031",
            zone,
            "00000100",
            datetime(2020, 1, 2),
            "",
            "",
            "",
        ]
    )
    wb.save(path)


def test_reads_current_workbook_layout_and_keeps_identifiers_as_text(tmp_path):
    path = tmp_path / "input.xlsx"
    make_book(path)
    data = read_workbook(path)
    assert not data.errors
    assert data.sheet_name == "РОСЗСП"
    assert data.rows[0].regno == "00000052"
    assert data.rows[0].include_order_number == "00000100"
    assert data.rows[0].include_order_date == datetime(2020, 1, 2)
    assert data.rows[0].licence_date_end_3 == datetime(2031, 1, 1)
    assert data.rows[0].licence_status_3 == "действующая"


def test_rejects_compound_numbering_zone(tmp_path):
    path = tmp_path / "input.xlsx"
    make_book(path, zone="843;855")
    data = read_workbook(path)
    assert any(issue.code == "geo_zone_format" for issue in data.errors)


def test_infers_missing_end_date_only_from_the_same_license_number(tmp_path):
    path = tmp_path / "input.xlsx"
    make_book(path)
    wb = load_workbook(path)
    ws = wb["РОСЗСП"]
    ws.cell(4, 8).value = "7777"
    copied = [ws.cell(4, column).value for column in range(1, len(HEADERS) + 1)]
    copied[2] = "00000053"
    copied[7] = "7777,"
    copied[8] = None
    ws.append(copied)
    wb.save(path)

    data = read_workbook(path)

    assert not data.errors
    assert data.rows[1].licence_number_1 == "7777"
    assert data.rows[1].licence_date_end_1 == datetime(2030, 1, 1)
    assert any(
        issue.code == "licence_end_inferred" and issue.excel_row == 5
        for issue in data.issues
    )


def test_reads_subject_fallback_columns_and_cleans_placeholder_values(tmp_path):
    path = tmp_path / "subject.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "РОСЗСП"
    ws.append(["Таблица миграции"])
    ws.append([])
    headers = HEADERS.copy()
    insert_at = headers.index("Полное наименование")
    headers[insert_at:insert_at] = ["ИНН", "Номер телефона", "Электронная почта", "ОПФ", "Код ОПФ"]
    address_at = headers.index("Место нахождения") + 1
    headers.insert(address_at, "Фактический/ Почтовый адрес")
    ws.append(headers)
    values = {
        "row_uuid": "uuid",
        "Идентификатор записи": "legacy",
        "Регистрационный номер": "52",
        "ОГРН": "1021600000001",
        "ИНН": "1655000000",
        "Номер телефона": "+7 843 000-00-00",
        "Электронная почта": "––",
        "ОПФ": "ПУБЛИЧНОЕ АКЦИОНЕРНОЕ ОБЩЕСТВО",
        "Код ОПФ": "12247",
        "Полное наименование": "ПАО ТЕСТОВЫЙ ОПЕРАТОР",
        "Сокращенное наименование": "ПАО ТЕСТ",
        "Место нахождения": "420000, г. Казань",
        "Фактический/ Почтовый адрес": "420000, г. Казань",
        "N1": "3004",
        "Дата начала и окончания лицензии_1": "01.01.2030",
        "Статус_1": "действующая",
        "N2": "3004",
        "Дата окончания лицензии_2": "01.06.2029",
        "Статус_2": "действующая",
        "N3": "9913",
        "Дата начала и окончания лицензии_3": "01.01.2031",
        "Статус_3": "действующая",
        "Зона нумерации": "843",
        "Номер приказа о включении": "100",
        "Дата приказа о включении": "02.01.2020",
    }
    ws.append([values.get(header, "") for header in headers])
    wb.save(path)

    row = read_workbook(path).rows[0]

    assert row.inn == "1655000000"
    assert row.phone == "+7 843 000-00-00"
    assert row.email == ""
    assert row.organizational_form_code == "12247"
    assert row.fact_address == "420000, г. Казань"
